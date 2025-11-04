#!/usr/bin/env python3
"""
Excel生成器
负责生成符合美团管家标准的菜品导入Excel文件
"""

import os
from datetime import datetime
from pathlib import Path
from typing import Dict, Any
import openpyxl
from openpyxl import load_workbook


class ExcelGenerator:
    """Excel文件生成器"""

    # 字段顺序 (32个字段)
    FIELD_ORDER = [
        "dish_name",              # 1. * 菜品名称
        "brand",                  # 2. * 所属品牌
        "category",               # 3. * 分类
        "spec",                   # 4. 规格
        "selling_price",          # 5. * 售卖价
        "member_price",           # 6. 会员价
        "estimated_cost",         # 7. 预估成本
        "barcode",                # 8. 条形码
        "pricing_method",         # 9. 计价方式
        "unit",                   # 10. 单位
        "numeric_memo",           # 11. 数字助记码
        "pinyin_memo",            # 12. 拼音助记码
        "shelf_life",             # 13. 保质期
        "shelf_life_unit",        # 14. 保质期单位
        "custom_field_1",         # 15. 自定义字段1
        "custom_field_2",         # 16. 自定义字段2
        "custom_field_3",         # 17. 自定义字段3
        "custom_field_4",         # 18. 自定义字段4
        "is_print",               # 19. 是否打印
        "selling_status",         # 20. 售卖状态
        "display_unified",        # 21. 点餐端展示-统一设置
        "allow_temp_price",       # 22. 收银端临时改价
        "allow_manual_discount",  # 23. 收银端手动打折
        "min_order_qty",          # 24. 起售份数
        "increment_qty",          # 25. 增量售卖数
        "combo_only",             # 26. 仅套餐售卖
        "dish_badge",             # 27. 菜品角标
        "description_tags",       # 28. 描述标签
        "order_tags",             # 29. 点餐标签
        "spicy_level",            # 30. 菜品辣度
        "description",            # 31. 菜品描述
        "detailed_description"    # 32. 菜品详细描述
    ]

    def __init__(self, template_path: str):
        """
        初始化Excel生成器

        Args:
            template_path: 模板文件路径
        """
        self.template_path = template_path

        if not os.path.exists(template_path):
            raise FileNotFoundError(f"模板文件不存在: {template_path}")

    def generate(
        self,
        data: Dict[str, Any],
        project_name: str,
        task_id: str
    ) -> str:
        """
        生成Excel文件

        Args:
            data: 菜品数据
            project_name: 项目名称
            task_id: 任务ID

        Returns:
            输出文件路径
        """
        # 1. 复制模板
        wb = load_workbook(self.template_path)
        ws = wb.active

        # 2. 填充数据到第2行 (第1行是表头)
        row_idx = 2
        for col_idx, field_name in enumerate(self.FIELD_ORDER, start=1):
            value = data.get(field_name)
            if value is not None:
                ws.cell(row=row_idx, column=col_idx, value=value)

        # 3. 保存到输出路径
        output_path = self._get_output_path(
            project_name=project_name,
            dish_name=data.get("dish_name"),
            task_id=task_id
        )

        # 创建输出目录
        output_dir = Path(output_path).parent
        output_dir.mkdir(parents=True, exist_ok=True)

        # 保存文件
        wb.save(output_path)

        return output_path

    def _get_output_path(
        self,
        project_name: str,
        dish_name: str,
        task_id: str
    ) -> str:
        """
        获取输出文件路径

        Args:
            project_name: 项目名称
            dish_name: 菜品名称
            task_id: 任务ID

        Returns:
            输出文件路径
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"菜品导入_{dish_name}_{timestamp}.xlsx"

        output_dir = Path("output") / project_name / "dish-creator"
        return str(output_dir / filename)


# 测试
if __name__ == "__main__":
    generator = ExcelGenerator(
        template_path="plugins/美团组/templates/菜品导入模板.xlsx"
    )

    test_data = {
        "dish_name": "麻辣小龙虾",
        "brand": "吼巷",
        "category": "招牌菜/海鲜",
        "selling_price": 88,
        "member_price": 78,
        "pricing_method": "按份销售",
        "unit": "份",
        "pinyin_memo": "MLXLX",
        "is_print": "是",
        "selling_status": "在售",
        "spicy_level": "重辣",
        "description": "精选优质小龙虾,秘制麻辣调料,鲜香麻辣"
    }

    output_path = generator.generate(
        data=test_data,
        project_name="测试项目",
        task_id="T001"
    )

    print(f"✅ Excel文件已生成: {output_path}")
